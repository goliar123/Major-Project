"""
Front Matter module for Major Project Phase-I Report (B12 Format).
"""
import os
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def add_front_matter(b):
    doc = b.doc
    LOGO_PATH = "/tmp/b12_media/word/media/image1.jpg"
    ACK_IMG_PATH = "/tmp/b12_media/word/media/image3.jpg"

    def set_cell_background(cell, hex_color):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # 1. TITLE PAGE
    b.p("A Major Project Phase-I Report On", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=10, space_after=8)
    b.p("AUTOMATED FAILURE ATTRIBUTION FOR LLM-BASED MULTI-AGENT SYSTEMS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, space_before=0, space_after=14, color=(30, 58, 138))
    b.p("Submitted in partial fulfillment of the requirements for the award of the\nBachelor of Technology\nIn\nDepartment of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, space_after=14)
    b.p("By", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=4)

    tbl_stud = doc.add_table(rows=5, cols=2)
    tbl_stud.alignment = WD_TABLE_ALIGNMENT.CENTER
    students_data = [
        ("Student Names", "Roll Numbers"),
        ("Rachapalli S Prabhaadhithya", "23241A05Q3"),
        ("Rathod Ajay", "23241A05Q4"),
        ("Noel Paul", "23241A05Q6"),
        ("Talagapu Niranjan", "23241A05R1")
    ]
    for r_idx, (name, rno) in enumerate(students_data):
        row = tbl_stud.rows[r_idx]
        row.cells[0].text = name
        row.cells[1].text = rno
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if r_idx == 0:
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[1].paragraphs[0].runs[0].bold = True
            set_cell_background(row.cells[0], "F1F5F9")
            set_cell_background(row.cells[1], "F1F5F9")
        set_cell_margins(row.cells[0], top=35, bottom=35, left=90, right=90)
        set_cell_margins(row.cells[1], top=35, bottom=35, left=90, right=90)

    b.p("", space_after=8)
    b.p("Under the Esteemed guidance of", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, space_before=4, space_after=2)
    b.p("Dr. D. Ushasree", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=2)
    b.p("Assistant Professor", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=10)

    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(4)
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.4))

    b.p("Department of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_before=4, space_after=2)
    b.p("GOKARAJU RANGARAJU INSTITUTE OF ENGINEERING AND TECHNOLOGY\n(Autonomous)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=2)
    b.p("Bachupally, Kukatpally, Hyderabad, Telangana, India, 500090\n2026-2027", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=0)

    b.page_break()

    # 2. CERTIFICATE
    if os.path.exists(LOGO_PATH):
        p_logo2 = doc.add_paragraph()
        p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo2.paragraph_format.space_before = Pt(0)
        p_logo2.paragraph_format.space_after = Pt(4)
        p_logo2.add_run().add_picture(LOGO_PATH, width=Inches(1.2))

    b.p("GOKARAJU RANGARAJU\nINSTITUTE OF ENGINEERING AND TECHNOLOGY\n(Autonomous)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=12)
    b.p("CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=12)

    cert_text = (
        "This is to certify that the major project phase-I work entitled “Automated Failure Attribution for "
        "LLM-Based Multi-Agent Systems” is submitted by Rachapalli S Prabhaadhithya (23241A05Q3), "
        "Rathod Ajay (23241A05Q4), Noel Paul (23241A05Q6), and Talagapu Niranjan (23241A05R1) in partial "
        "fulfillment of the award of a degree in BACHELOR OF TECHNOLOGY in Computer Science and Engineering "
        "during the academic year 2026-2027."
    )
    p_cert = b.p(cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_before=10, space_after=26)
    p_cert.paragraph_format.line_spacing = 1.3

    tbl_sig = doc.add_table(rows=3, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.rows[0].cells[0].text = "GUIDE"
    tbl_sig.rows[0].cells[1].text = "HEAD OF THE DEPARTMENT"
    tbl_sig.rows[1].cells[0].text = "Dr. D. Ushasree"
    tbl_sig.rows[1].cells[1].text = "Dr. B. SANKARA BABU"
    tbl_sig.rows[2].cells[0].text = "Assistant Professor\nDept. of CSE"
    tbl_sig.rows[2].cells[1].text = "Professor\nDept. of CSE"

    for r in tbl_sig.rows:
        r.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r.cells[0].paragraphs[0].runs[0].bold = True
        r.cells[1].paragraphs[0].runs[0].bold = True

    b.p("", space_after=35)
    b.p("EXTERNAL EXAMINER", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_before=16)
    b.p("_____________________", align=WD_ALIGN_PARAGRAPH.LEFT, size=12)

    b.page_break()

    # 3. ACKNOWLEDGEMENT
    if os.path.exists(ACK_IMG_PATH):
        p_ack = doc.add_paragraph()
        p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ack.paragraph_format.space_before = Pt(4)
        p_ack.paragraph_format.space_after = Pt(10)
        p_ack.add_run().add_picture(ACK_IMG_PATH, width=Inches(3.3))
    else:
        b.p("ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=10)

    ack_text = (
        "Many people helped us directly and indirectly to complete our project successfully. We would like to take "
        "this opportunity to thank one and all. First, we wish to express our deep gratitude to our guide Dr. D. Ushasree, "
        "Assistant Professor, Department of CSE for her support in the completion of our project work Phase-I successfully "
        "and for all the time to time guidance provided. The process made us to learn a lot practically and we are very "
        "thankful to the final year Project Coordinator Dr. N. Krishna Chaithanya, Asst. Prof., CSE and our class project "
        "coordinator for their unwavering support in providing schedules, formats, rubrics and arranging the seminars at "
        "regular intervals. Our Sincere thanks to Project Review Committee member Dr. K. Prasanna Lakshmi, Professor & Dean "
        "Academic Affairs, for their in-depth analysis during evaluations of seminars that helped us in improvising the quality "
        "of our work. We wish to express our honest and sincere thanks to Dr. B. Sankara Babu, HOD, Department of CSE, to our "
        "Principal Dr. J. Praveen and to our Director Dr. Jandhyala N Murthy for providing all the facilities required to "
        "complete our major project Phase-I. We would like to thank all our faculty and friends for their help and constructive "
        "criticism during the completion of this phase. Finally, we are very much indebted to our parents for their moral support "
        "and encouragement to achieve goals."
    )
    p_ack_b = b.p(ack_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_before=6, space_after=20)
    p_ack_b.paragraph_format.line_spacing = 1.25

    b.p("Student Names\t(Roll Numbers)", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11, space_after=4)
    for name, rno in [("Rachapalli S Prabhaadhithya", "23241A05Q3"), ("Rathod Ajay", "23241A05Q4"), ("Noel Paul", "23241A05Q6"), ("Talagapu Niranjan", "23241A05R1")]:
        b.p(f"{name}\t{rno}", align=WD_ALIGN_PARAGRAPH.RIGHT, size=11, space_after=2)

    b.page_break()

    # 4. DECLARATION
    b.p("DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=10, space_after=12)
    decl_text = (
        "We hereby declare that the major project phase-I entitled “Automated Failure Attribution for LLM-Based Multi-Agent "
        "Systems” is the work done during the period from 2026-2027 and is submitted in the partial fulfillment of the requirements "
        "for the award of the degree of Bachelor of Technology in Computer Science and Engineering from Gokaraju Rangaraju "
        "Institute of Engineering and Technology (Autonomous). The results embodied in this phase-I of project have not been "
        "submitted to any other University or Institution for the award of any degree or diploma."
    )
    p_dec = b.p(decl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_before=8, space_after=26)
    p_dec.paragraph_format.line_spacing = 1.3

    b.p("Students Names\t(Roll Numbers)", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11, space_after=4)
    for name, rno in [("Rachapalli S Prabhaadhithya", "23241A05Q3"), ("Rathod Ajay", "23241A05Q4"), ("Noel Paul", "23241A05Q6"), ("Talagapu Niranjan", "23241A05R1")]:
        b.p(f"{name}\t{rno}", align=WD_ALIGN_PARAGRAPH.RIGHT, size=11, space_after=2)

    b.page_break()

    # 5. TABLE OF CONTENTS
    b.p("Table of Contents", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=8)
    
    toc_data = [
        ("Chapter", "TITLE", "Page No."),
        ("", "Abstract", "1"),
        ("1", "Introduction", "2-3"),
        ("2", "Literature Survey", "4-8"),
        ("", "2.1 Introduction", "4"),
        ("", "2.2 Related Work", "4-7"),
        ("", "2.2.1 Multi-Agent Orchestration Frameworks & Failure Dynamics", "4"),
        ("", "2.2.2 Causal Graph Modeling & Transitive Dependency Analysis", "5"),
        ("", "2.2.3 LLM-as-a-Judge and Semantic Verification", "5-6"),
        ("", "2.2.4 Benchmarking Multi-Agent Failure Attribution (The Who&When Paradigm)", "6"),
        ("", "2.2.5 Software Observability & Distributed Tracing Adaptations", "6-7"),
        ("", "2.2.6 Comparative Analysis of Diagnostic Paradigms (Table 2.1)", "7"),
        ("", "2.3 Gaps Identified", "7-8"),
        ("", "2.4 Problem Statement", "8"),
        ("", "2.5 Proposed Solution", "8"),
        ("", "2.6 Summary", "8"),
        ("3", "System Requirements Specification", "9-44"),
        ("", "3.1 Introduction", "9"),
        ("", "3.1.1 Purpose", "9"),
        ("", "3.1.2 Definitions, Acronyms and Abbreviations", "10-13"),
        ("", "3.1.3 Scope", "13-14"),
        ("", "3.1.4 References", "14"),
        ("", "3.1.5 Overview", "14"),
        ("", "3.2 General Description", "14-22"),
        ("", "3.2.1 Product Perspective", "14-15"),
        ("", "3.2.2 Product Functions", "15-17"),
        ("", "3.2.3 User Characteristics", "18-19"),
        ("", "3.2.4 General Constraints", "19-20"),
        ("", "3.2.5 Assumptions and Dependencies", "21-22"),
        ("", "3.3 Specific Requirements", "22-35"),
        ("", "3.3.1 Functional Requirements", "22-26"),
        ("", "3.3.2 Non-Functional Requirements", "26-29"),
        ("", "3.3.3 User Interface Requirements", "29-31"),
        ("", "3.3.4 Metadata and Schema of Database (Tables 3.5 - 3.11)", "31-33"),
        ("", "3.3.5 Hardware Requirements (Tables 3.3.5.1 - 3.3.5.2)", "34-35"),
        ("", "3.4 Architecture and UML Diagrams", "35-44"),
        ("", "3.4.1 Architecture Diagram", "35-37"),
        ("", "3.4.2 Use Case Diagram", "37-38"),
        ("", "3.4.3 Class Diagram", "38-40"),
        ("", "3.4.4 Sequence Diagram", "40-41"),
        ("", "3.4.5 Data Flow Diagram - Level 0", "41-42"),
        ("", "3.4.6 Data Flow Diagram - Level 1", "42-43"),
        ("", "3.4.7 Entity Relationship Diagram", "43-44"),
        ("", "3.5 Summary", "44"),
        ("4", "Methodology", "45-54"),
        ("", "4.1 Modules (Table 4.1)", "45-46"),
        ("", "4.2 Methodology Proposed For Solution Implementation", "46-53"),
        ("", "4.2.1 Core Algorithms and Python Implementation Listings (Listings 4.1-4.5)", "47-51"),
        ("", "4.2.2 25+ Failure Pattern Taxonomy & MockJudge Calibrated Scores (Table 4.2)", "51-52"),
        ("", "4.2.3 Empirical Benchmarking on ICML 2025 Who&When Callset (Table 4.3)", "52-53"),
        ("", "4.3 Summary", "54"),
        ("5", "Conclusion and Future Scope", "55-58"),
        ("", "5.1 Conclusion", "55-56"),
        ("", "5.2 Future Scope", "56-58"),
        ("6", "References", "59-60"),
        ("", "Appendix: Research Article Draft & Submission Record", "61-64")
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_data), cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, (c1, c2, c3) in enumerate(toc_data):
        row = tbl_toc.rows[r_idx]
        row.cells[0].text = c1
        row.cells[1].text = c2
        row.cells[2].text = c3
        if r_idx == 0:
            set_cell_background(row.cells[0], "E2E8F0")
            set_cell_background(row.cells[1], "E2E8F0")
            set_cell_background(row.cells[2], "E2E8F0")
            for c in row.cells:
                c.paragraphs[0].runs[0].bold = True
        set_cell_margins(row.cells[0], top=18, bottom=18, left=35, right=35)
        set_cell_margins(row.cells[1], top=18, bottom=18, left=35, right=35)
        set_cell_margins(row.cells[2], top=18, bottom=18, left=35, right=35)

    b.p("", space_after=8)
    b.p("LIST OF FIGURES", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_before=8, space_after=4)
    lof_data = [
        ("Fig No", "Fig Title", "Page No"),
        ("3.1", "Architecture Diagram of Automated Failure Attribution Engine", "36"),
        ("3.2", "Use Case Diagram of Failure Attribution System", "38"),
        ("3.3", "Class Diagram of Failure Attribution System", "40"),
        ("3.4", "Sequence Diagram of Multi-Agent Attribution Workflow", "41"),
        ("3.5", "Data Flow Level - 0 Diagram", "42"),
        ("3.6", "Data Flow Level - 1 Diagram", "43"),
        ("3.7", "Entity Relationship Diagram of Diagnostic Schema", "44")
    ]
    tbl_lof = doc.add_table(rows=len(lof_data), cols=3)
    tbl_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, (c1, c2, c3) in enumerate(lof_data):
        row = tbl_lof.rows[r_idx]
        row.cells[0].text = c1
        row.cells[1].text = c2
        row.cells[2].text = c3
        if r_idx == 0:
            set_cell_background(row.cells[0], "E2E8F0")
            set_cell_background(row.cells[1], "E2E8F0")
            set_cell_background(row.cells[2], "E2E8F0")
            for c in row.cells:
                c.paragraphs[0].runs[0].bold = True
        set_cell_margins(row.cells[0], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[1], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[2], top=20, bottom=20, left=40, right=40)

    b.p("", space_after=8)
    b.p("LIST OF TABLES", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_before=8, space_after=4)
    lot_data = [
        ("Table No", "Table Name", "Page No"),
        ("2.1", "Comparative Matrix of Failure Attribution Paradigms", "7"),
        ("3.1.3", "Definitions, Acronyms and Abbreviations", "10-13"),
        ("3.2.5", "Assumptions and Dependencies", "22"),
        ("3.5", "multi_agent_traces Table Schema", "31"),
        ("3.6", "execution_steps Table Schema", "31"),
        ("3.7", "step_evaluations Table Schema", "32"),
        ("3.8", "dependency_edges Table Schema", "32"),
        ("3.9", "prompt_cache Table Schema", "33"),
        ("3.10", "attribution_verdicts Table Schema", "33"),
        ("3.11", "benchmark_results Table Schema", "33"),
        ("3.3.5.1", "Minimum Hardware Requirements", "34"),
        ("3.3.5.2", "Recommended Hardware Requirements", "34-35"),
        ("4.1", "Modules of Failure Attribution System", "45-46"),
        ("4.2", "Comprehensive 25+ Failure Pattern Taxonomy & MockJudge Scores", "51-52"),
        ("4.3", "Empirical Accuracy and Latency Benchmark on Who&When Callset", "52-53")
    ]
    tbl_lot = doc.add_table(rows=len(lot_data), cols=3)
    tbl_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, (c1, c2, c3) in enumerate(lot_data):
        row = tbl_lot.rows[r_idx]
        row.cells[0].text = c1
        row.cells[1].text = c2
        row.cells[2].text = c3
        if r_idx == 0:
            set_cell_background(row.cells[0], "E2E8F0")
            set_cell_background(row.cells[1], "E2E8F0")
            set_cell_background(row.cells[2], "E2E8F0")
            for c in row.cells:
                c.paragraphs[0].runs[0].bold = True
        set_cell_margins(row.cells[0], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[1], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[2], top=20, bottom=20, left=40, right=40)

    b.p("", space_after=8)
    b.p("LIST OF CODE LISTINGS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_before=8, space_after=4)
    lol_data = [
        ("Listing No", "Listing Title", "Page No"),
        ("4.1", "Asymmetric Context Token Budgeting Implementation", "48"),
        ("4.2", "Deterministic SHA256 Hashing & Two-Tier Cache Architecture", "49"),
        ("4.3", "Dependency DAG Construction & Entity Linking", "50"),
        ("4.4", "Transitive Ancestor Closure Computation", "50"),
        ("4.5", "Topological Blame Resolution & Clean-Parent Invariant", "51")
    ]
    tbl_lol = doc.add_table(rows=len(lol_data), cols=3)
    tbl_lol.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, (c1, c2, c3) in enumerate(lol_data):
        row = tbl_lol.rows[r_idx]
        row.cells[0].text = c1
        row.cells[1].text = c2
        row.cells[2].text = c3
        if r_idx == 0:
            set_cell_background(row.cells[0], "E2E8F0")
            set_cell_background(row.cells[1], "E2E8F0")
            set_cell_background(row.cells[2], "E2E8F0")
            for c in row.cells:
                c.paragraphs[0].runs[0].bold = True
        set_cell_margins(row.cells[0], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[1], top=20, bottom=20, left=40, right=40)
        set_cell_margins(row.cells[2], top=20, bottom=20, left=40, right=40)

    b.page_break()

    # 6. ABSTRACT
    b.h1("ABSTRACT")
    b.p(
        "Communicative multi-agent systems powered by Large Language Models (LLMs)—including Microsoft AutoGen, "
        "MetaGPT, and ChatDev—enable autonomous collaborative problem-solving across complex engineering domains. "
        "However, when an agent in an execution cascade introduces an erroneous premise, invalid tool parameter, or "
        "unfulfilled schema constraint, the defect propagates silently across subsequent conversational turns, ultimately "
        "culminating in a catastrophic crash at an end-of-chain agent. Conventional monitoring frameworks and monolithic "
        "LLM-as-a-judge evaluators suffer from severe recency bias, consistently committing 'Symptom Misattribution' by "
        "blaming the terminal crashing node rather than the silent upstream originator, resulting in less than 15% step-level "
        "accuracy on standardized benchmarks.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=8
    )
    b.p(
        "To resolve this fundamental diagnostic challenge, this project introduces a Neuro-Symbolic Causal Graph Engine "
        "that cleanly decouples localized neural semantic evaluation from global symbolic graph causality. Discrete execution "
        "turns are evaluated by a locally hosted Qwen2.5:7b model via Ollama in structured JSON mode, accelerated by a "
        "deterministic SHA256 prompt cache backed by SQLite that eliminates 100% of redundant inference. Concurrently, a "
        "Directed Acyclic Graph (DAG) is constructed using NetworkX, capturing sequential conversational flow (confidence 0.85) "
        "and cross-turn variable dependencies (confidence 0.95) extracted via regex entity matching. A Two-Phase Backward "
        "Causal Traversal algorithm traverses backward from the manifest failure symptom, computes the Transitive Ancestor "
        "Closure, and enforces a Clean-Parent Invariant to mathematically isolate the true originating root cause. "
        "Evaluated on the standardized ICML 2025 Who&When benchmark comprising 184 traces across six diverse environments, "
        "the engine achieves 85.3% agent accuracy (Who) and 78.3% step accuracy (When) while cutting inference latency by 91.2%, "
        "providing essential, production-grade observability infrastructure for reliable multi-agent systems.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=12
    )
    b.page_break()
